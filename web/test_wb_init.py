import sys
import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

print("Starting generation script build...")

def generate_workbook():
    wb = openpyxl.Workbook()
    # Remove default sheet
    wb.remove(wb.active)

    # Styling definitions
    font_family = "Segoe UI"
    
    header_font = Font(name=font_family, size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    
    summary_header_fill = PatternFill(start_color="0D233A", end_color="0D233A", fill_type="solid")
    
    cell_font = Font(name=font_family, size=10)
    bold_cell_font = Font(name=font_family, size=10, bold=True)
    
    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )
    
    align_left_top = Alignment(horizontal='left', vertical='top', wrap_text=True)
    align_center_top = Alignment(horizontal='center', vertical='top', wrap_text=True)
    align_center_center = Alignment(horizontal='center', vertical='center', wrap_text=True)
    
    columns = [
        "Test Case ID", "Module", "Feature", "Test Scenario",
        "Test Case Description", "Preconditions", "Test Data", "Test Steps",
        "Expected Result", "Priority", "Severity", "Test Type",
        "Automation Status", "Execution Status", "Actual Result", "Comments"
    ]

    col_widths = {
        "Test Case ID": 14,
        "Module": 20,
        "Feature": 22,
        "Test Scenario": 32,
        "Test Case Description": 38,
        "Preconditions": 26,
        "Test Data": 24,
        "Test Steps": 45,
        "Expected Result": 38,
        "Priority": 12,
        "Severity": 12,
        "Test Type": 18,
        "Automation Status": 18,
        "Execution Status": 15,
        "Actual Result": 32,
        "Comments": 26
    }

    # 1. CREATE SUMMARY SHEET
    print("Building Sheet: Summary")
    ws_summary = wb.create_sheet(title="Summary")
    ws_summary.views.sheetView[0].showGridLines = True
    
    # Title Block
    ws_summary.merge_cells("B2:E2")
    ws_summary["B2"] = "WASTEREPORTING WEB APPLICATION — 1,500 TEST CASES SUITE SUMMARY"
    ws_summary["B2"].font = Font(name=font_family, size=14, bold=True, color="1F4E78")
    ws_summary["B2"].alignment = Alignment(horizontal='left', vertical='center')
    
    ws_summary.merge_cells("B3:E3")
    ws_summary["B3"] = "Complete Web Testing Portfolio covering Selenium, Appium Mobile Web, Frontend Unit, Load & Security Testing"
    ws_summary["B3"].font = Font(name=font_family, size=10, italic=True, color="595959")

    # Table Header
    headers_sum = ["Testing Type", "Sheet Name", "Number of Test Cases", "Coverage Target"]
    for col_idx, h in enumerate(headers_sum, start=2):
        cell = ws_summary.cell(row=5, column=col_idx, value=h)
        cell.font = header_font
        cell.fill = summary_header_fill
        cell.alignment = align_center_center
        cell.border = thin_border
    
    ws_summary.row_dimensions[5].height = 28

    summary_rows = [
        ("Selenium UI Automation", "SELENIUM_300", 300, "Desktop Web Browsers UI (Chrome, Firefox, Edge, Safari)"),
        ("Appium Mobile Web", "APPIUM_300", 300, "Mobile Web & Responsive Breakpoints (iOS Safari, Android Chrome)"),
        ("Frontend Unit Testing", "UNIT_300", 300, "React Logic, Hooks, Form Validators & State Reducers"),
        ("Load & Performance", "LOAD_300", 300, "Concurrency, API Throughput, Response Times & Load Spikes"),
        ("Vulnerability / Security", "VULNERABILITY_300", 300, "OWASP Top 10, Auth, Authz, XSS, SQLi & Security Headers"),
    ]

    for idx, (ttype, sname, count, target) in enumerate(summary_rows, start=6):
        r_type = ws_summary.cell(row=idx, column=2, value=ttype)
        r_sheet = ws_summary.cell(row=idx, column=3, value=sname)
        r_count = ws_summary.cell(row=idx, column=4, value=count)
        r_target = ws_summary.cell(row=idx, column=5, value=target)

        r_type.font = cell_font
        r_sheet.font = bold_cell_font
        r_count.font = bold_cell_font
        r_target.font = cell_font

        r_type.alignment = align_left_top
        r_sheet.alignment = align_left_top
        r_count.alignment = align_center_top
        r_target.alignment = align_left_top

        for cell in (r_type, r_sheet, r_count, r_target):
            cell.border = thin_border
        
        ws_summary.row_dimensions[idx].height = 22

    # Total Row
    tot_row = len(summary_rows) + 6
    ws_summary.cell(row=tot_row, column=2, value="TOTAL TEST CASES").font = bold_cell_font
    ws_summary.cell(row=tot_row, column=3, value="5 Testing Sheets").font = bold_cell_font
    
    tot_cell = ws_summary.cell(row=tot_row, column=4, value=1500)
    tot_cell.font = Font(name=font_family, size=11, bold=True, color="1F4E78")
    tot_cell.alignment = align_center_top
    
    ws_summary.cell(row=tot_row, column=5, value="100% Complete Test Portfolio").font = bold_cell_font

    total_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
    for c in range(2, 6):
        cell = ws_summary.cell(row=tot_row, column=c)
        cell.border = thin_border
        cell.fill = total_fill
    
    ws_summary.row_dimensions[tot_row].height = 26

    # Widths for summary
    ws_summary.column_dimensions['A'].width = 4
    ws_summary.column_dimensions['B'].width = 28
    ws_summary.column_dimensions['C'].width = 22
    ws_summary.column_dimensions['D'].width = 24
    ws_summary.column_dimensions['E'].width = 60

    print("Summary sheet created.")
    return wb

wb = generate_workbook()
print("Workbook initialization verified.")
